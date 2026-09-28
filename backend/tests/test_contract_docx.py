import unittest

from docx import Document
from app.services.contract_docx import (
    _replace_across_runs,
    _tighten_legal_section,
    parse_service_address,
)


class ContractDocxTestCase(unittest.TestCase):
    def test_structured_address_from_manual_service_address(self) -> None:
        address = parse_service_address(
            "Calle Fiesta 2706-B, Int. 2, Fraccionamiento Miravalle, "
            "C.P. 88715, Reynosa, Tamaulipas, México"
        )

        self.assertEqual(address.street, "Fiesta")
        self.assertEqual(address.exterior, "2706-B")
        self.assertEqual(address.interior, "2")
        self.assertEqual(address.settlement, "Miravalle")
        self.assertEqual(address.city, "Reynosa")
        self.assertEqual(address.state, "Tamaulipas")
        self.assertEqual(address.postal_code, "88715")

    def test_replacement_preserves_the_formatting_run(self) -> None:
        paragraph = Document().add_paragraph()
        first = paragraph.add_run("Jose")
        first.bold = True
        paragraph.add_run("fina")

        _replace_across_runs(paragraph, "Josefina", "Marcelina")

        self.assertEqual(paragraph.text, "Marcelina")
        self.assertTrue(paragraph.runs[0].bold)

    def test_equal_replacement_returns_without_looping(self) -> None:
        paragraph = Document().add_paragraph("Ubiquiti")

        _replace_across_runs(paragraph, "Ubiquiti", "Ubiquiti")

        self.assertEqual(paragraph.text, "Ubiquiti")

    def test_replacement_may_contain_original_text(self) -> None:
        paragraph = Document().add_paragraph("Ubiquiti")

        _replace_across_runs(paragraph, "Ubiquiti", "Ubiquiti Networks")

        self.assertEqual(paragraph.text, "Ubiquiti Networks")

    def test_contract_annex_starts_on_a_new_page(self) -> None:
        document = Document()
        document.add_paragraph("CONTRATO DE PRESTACIÓN DE SERVICIO")
        annex = document.add_paragraph("ANEXO 1 DEL CONTRATO")

        _tighten_legal_section(document)

        self.assertTrue(annex.paragraph_format.page_break_before)


if __name__ == "__main__":
    unittest.main()
