"""Render AMR contracts from the editable Word master without overlaying text."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt


@dataclass(frozen=True)
class ContractAddress:
    street: str
    exterior: str
    interior: str
    settlement: str
    city: str
    state: str
    postal_code: str

    @property
    def installation_line(self) -> str:
        number = f" {self.exterior}" if self.exterior else ""
        interior = f" Int. {self.interior}" if self.interior else ""
        settlement = f", {self.settlement}" if self.settlement else ""
        return f"{self.street}{number}{interior}{settlement}".strip(", ")


def parse_service_address(value: str) -> ContractAddress:
    """Split the canonical address written by Aether into the Word form slots."""
    parts = [part.strip() for part in value.split(",") if part.strip()]
    postal_code = ""
    cleaned_parts: list[str] = []
    for part in parts:
        postal_match = re.search(r"(?:C\.?\s*P\.?\s*)?(\d{5})\b", part, re.IGNORECASE)
        if postal_match:
            postal_code = postal_match.group(1)
            part = (part[: postal_match.start()] + part[postal_match.end() :]).strip(" .-")
        if part:
            cleaned_parts.append(part)
    parts = cleaned_parts

    if parts and parts[-1].casefold() in {"méxico", "mexico"}:
        parts.pop()
    state = parts.pop() if len(parts) > 2 else "Tamaulipas"
    city = parts.pop() if len(parts) > 1 else "Reynosa"

    first = parts.pop(0) if parts else value
    first = re.sub(r"^calle\s+", "", first, flags=re.IGNORECASE).strip()
    exterior_match = re.match(r"^(.*?)(?:\s+#?)(\d+[\w-]*)$", first)
    street = exterior_match.group(1).strip() if exterior_match else first
    exterior = exterior_match.group(2).strip() if exterior_match else ""

    interior = ""
    if parts and re.match(r"^(?:int\.?|interior)\s*", parts[0], re.IGNORECASE):
        interior = re.sub(r"^(?:int\.?|interior)\s*", "", parts.pop(0), flags=re.IGNORECASE)
    settlement = parts[0] if parts else ""
    settlement = re.sub(
        r"^(?:fraccionamiento|fracc\.?|colonia|col\.?)\s+",
        "",
        settlement,
        flags=re.IGNORECASE,
    ).strip()
    return ContractAddress(
        street=street,
        exterior=exterior,
        interior=interior,
        settlement=settlement,
        city=city,
        state=state,
        postal_code=postal_code,
    )


def _replace_across_runs(paragraph, source: str, replacement: str) -> None:
    """Replace text split across Word runs while retaining the original formatting."""
    if not source or source == replacement:
        return
    search_from = 0
    while True:
        runs = paragraph.runs
        full_text = "".join(run.text for run in runs)
        start = full_text.find(source, search_from)
        if start < 0:
            break
        end = start + len(source)
        cursor = 0
        start_run = end_run = 0
        start_offset = end_offset = 0
        for index, run in enumerate(runs):
            next_cursor = cursor + len(run.text)
            if cursor <= start < next_cursor:
                start_run = index
                start_offset = start - cursor
            if cursor < end <= next_cursor:
                end_run = index
                end_offset = end - cursor
                break
            cursor = next_cursor

        prefix = runs[start_run].text[:start_offset]
        suffix = runs[end_run].text[end_offset:]
        runs[start_run].text = prefix + replacement
        for index in range(start_run + 1, end_run):
            runs[index].text = ""
        if end_run == start_run:
            runs[start_run].text += suffix
        else:
            runs[end_run].text = suffix
        search_from = start + len(replacement)


def _replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    for source in sorted(replacements, key=len, reverse=True):
        _replace_across_runs(paragraph, source, replacements[source])


def _replace_story(paragraphs, tables, replacements: dict[str, str]) -> None:
    for paragraph in paragraphs:
        _replace_in_paragraph(paragraph, replacements)
    for table in tables:
        for row in table.rows:
            seen_cells: set[int] = set()
            for cell in row.cells:
                cell_key = id(cell._tc)
                if cell_key in seen_cells:
                    continue
                seen_cells.add(cell_key)
                _replace_story(cell.paragraphs, cell.tables, replacements)


def _set_empty_cell(cell, value: str) -> None:
    if not value:
        return
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
    else:
        run = paragraph.add_run(value)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.bold = True


def _set_cell_font_size(cell, points: float) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(points)


def _set_paragraph_text(paragraph, text: str, points: float) -> None:
    """Reuse the first formatted run and discard spacing that Word may reflow."""
    if not paragraph.runs:
        paragraph.add_run()
    paragraph.runs[0].text = text
    paragraph.runs[0].font.size = Pt(points)
    for run in paragraph.runs[1:]:
        run.text = ""


def _align_address_labels(cell) -> None:
    paragraph = cell.paragraphs[0]
    labels = ("Calle", "#Ext.", "#Int.", "Colonia", "Alcaldía/Municipio", "Estado", "C.P.")
    positions = (90, 198, 225, 286, 385, 462, 520)
    paragraph.paragraph_format.tab_stops.clear_all()
    for position in positions:
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Pt(position),
            WD_TAB_ALIGNMENT.CENTER,
        )
    _set_paragraph_text(paragraph, "\t" + "\t".join(labels), 7)


def _tighten_legal_section(document) -> None:
    """Match Word's seven-page pagination when LibreOffice performs the export."""
    inside_legal_text = False
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("ANEXO 1 DEL CONTRATO"):
            # LibreOffice ignores the Word continuous-section boundary here.
            # A paragraph-level break keeps the annex and promissory note on
            # their intended final sheet without producing an empty page.
            paragraph.paragraph_format.page_break_before = True
            break
        if paragraph.text.startswith("CONTRATO DE PRESTACIÓN DE SERVICIO"):
            inside_legal_text = True
        if inside_legal_text:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            for run in paragraph.runs:
                run.font.size = Pt(8.75)


def render_contract(template: Path, destination: Path, values: dict[str, str]) -> None:
    document = Document(template)
    address = ContractAddress(
        street=values["street"],
        exterior=values["exterior"],
        interior=values["interior"],
        settlement=values["settlement"],
        city=values["city"],
        state=values["state"],
        postal_code=values["postal_code"],
    )
    number = values["contract_number"]
    replacements = {
        "Calle Barra Arida 136 Puerta Sur, Reynosa Tamaulipas": values["address"],
        "Calle Barra arida 136 Puerta Sur": address.installation_line,
        "Josefina Ramirez Treviño": values["customer_name"],
        "11 días de mes de Febrero del año 2018": values["annex_date"],
        "11 de Febrero del 2018": values["contract_date_long"],
        "11/02/2018": values["contract_date"],
        "CONTRATO CON NÚMERO 132": f"CONTRATO CON NÚMERO {number}",
        "CONTRATO CON NÚMERO 250": f"CONTRATO CON NÚMERO {number}",
        "NÚMERO 132": f"NÚMERO {number}",
        "NÚMERO 250": f"NÚMERO {number}",
        "Numero 132": f"Numero {number}",
        "Numero 250": f"Numero {number}",
        "Hasta 10 megas de velocidad": values["plan_name"],
        "$ 350, M.N": values["monthly_price"],
        "11 de cada mes": values["payment_day"],
        "8992088778": values["phone"],
        "Josefina": values["given_names"],
        "Ramirez": values["paternal_surname"],
        "Treviño": values["maternal_surname"],
        "Barra Arida": address.street,
        "Puerta Sur": address.settlement,
        "88734": address.postal_code,
    }
    _replace_story(document.paragraphs, document.tables, replacements)
    for section in document.sections:
        _replace_story(section.header.paragraphs, section.header.tables, replacements)
        _replace_story(section.footer.paragraphs, section.footer.tables, replacements)

    # Values such as "136" are not safe global tokens; use their table slots.
    subscriber_table = document.tables[0]
    _set_cell_font_size(subscriber_table.cell(1, 2), 8.5)
    _set_cell_font_size(subscriber_table.cell(1, 5), 8.5)
    _set_cell_font_size(subscriber_table.cell(1, 18), 7)
    _replace_story(subscriber_table.cell(4, 6).paragraphs, [], {"136": address.exterior})
    _set_cell_font_size(subscriber_table.cell(4, 6), 7)
    _set_empty_cell(subscriber_table.cell(4, 7), address.interior)
    _replace_story(subscriber_table.cell(4, 18).paragraphs, [], {"Reynosa": address.city})
    _replace_story(subscriber_table.cell(4, 20).paragraphs, [], {"Tamaulipas": address.state})
    _set_cell_font_size(subscriber_table.cell(4, 18), 7.5)
    _set_cell_font_size(subscriber_table.cell(4, 20), 7.5)
    _align_address_labels(subscriber_table.cell(5, 0))
    phone_label = subscriber_table.cell(6, 0)
    _set_paragraph_text(phone_label.paragraphs[0], "TELÉFONO Fijo     Móvil X", 7)
    phone_label.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    pricing_table = document.tables[1]
    _set_cell_font_size(pricing_table.cell(2, 2), 7)

    equipment_table = document.tables[2]
    _replace_story(equipment_table.cell(1, 1).paragraphs, [], {"Ubiquiti": values["equipment_brand"]})
    _replace_story(equipment_table.cell(2, 1).paragraphs, [], {"Lbe 5ac gen2": values["equipment_model"]})
    _replace_story(equipment_table.cell(3, 1).paragraphs, [], {"n/a": values["equipment_serial"]})
    _replace_story(equipment_table.cell(4, 1).paragraphs, [], {"AMR132": values["service_code"]})
    installation_table = document.tables[3]
    _replace_story(installation_table.cell(3, 1).paragraphs, [], {"$1300": values["installation_cost"]})
    _set_cell_font_size(installation_table.cell(2, 0), 7)
    _set_cell_font_size(installation_table.cell(2, 2), 7)
    _set_cell_font_size(installation_table.cell(3, 0), 7)

    details_table = document.tables[6]
    for row, column in ((2, 21), (6, 21), (9, 0), (14, 0)):
        _set_cell_font_size(details_table.cell(row, column), 7)

    _tighten_legal_section(document)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="aether-contract-") as work:
        docx_path = Path(work) / "contrato.docx"
        profile_path = Path(work) / "libreoffice-profile"
        document.save(docx_path)
        result = subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation={profile_path.as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                work,
                str(docx_path),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=90,
        )
        generated = Path(work) / "contrato.pdf"
        if result.returncode != 0 or not generated.is_file():
            diagnostic = (result.stderr or result.stdout or "sin detalle").strip()
            raise RuntimeError(f"No fue posible convertir el contrato Word a PDF: {diagnostic}")
        destination.write_bytes(generated.read_bytes())
